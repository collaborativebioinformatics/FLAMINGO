"""Build four visual slides from committed FLAMINGO results.
Requires matplotlib, numpy, Pillow, python-pptx. Run from any directory.
"""
from pathlib import Path
import csv, json, math
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Circle, Rectangle, Arc
from matplotlib.backends.backend_pdf import PdfPages
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

OUT=Path(__file__).resolve().parent
ROOT=OUT.parent
BG='#19161d'; PANEL='#26212c'; PINK='#f078a5'; WHITE='#fff5f8'; MUTED='#b4a5b8'; TEAL='#7cdecf'; EDGE='#49384d'; GOLD='#edca83'
plt.rcParams.update({'font.family':'DejaVu Sans','text.color':WHITE,'font.size':16,'svg.fonttype':'none'})
figures=[]; notes=[]
def text(ax,x,y,s,size=18,color=WHITE,ha='left',weight='normal',**kw):
    return ax.text(x,y,s,fontsize=size,color=color,ha=ha,va='center',weight=weight,**kw)
def line(ax,xs,ys,c=PINK,lw=2,**kw): ax.plot(xs,ys,color=c,lw=lw,solid_capstyle='round',**kw)
def card(ax,x,y,w,h,color=PANEL,edge=EDGE):
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.02,rounding_size=0.2',facecolor=color,edgecolor=edge,lw=1.2))
def arrow(ax,a,b,c=PINK,lw=2):
    ax.annotate('',xy=b,xytext=a,arrowprops={'arrowstyle':'-|>','color':c,'lw':lw,'mutation_scale':20})
def base(num,title):
    f=plt.figure(figsize=(16,9),facecolor=BG)
    ax=f.add_axes([0,0,1,1]); ax.set(xlim=(0,16),ylim=(0,9)); ax.axis('off')
    text(ax,.65,8.35,title,30,weight='bold')
    text(ax,15.35,8.35,'FLAMINGO',15,PINK,ha='right',weight='bold')
    line(ax,[.65,15.35],[.55,.55],EDGE,1)
    text(ax,.65,.28,'FEDERATED MENDELIAN RANDOMIZATION',9,MUTED)
    text(ax,15.35,.28,f'{num:02d} / 04',11,MUTED,ha='right')
    figures.append(f); return f,ax

def bank(ax,x,y,s=1):
    # A site containing person-level records; records remain inside the site.
    card(ax,x-.55*s,y-.63*s,1.1*s,1.26*s)
    line(ax,[x-.4*s,x,x+.4*s],[y+.3*s,y+.53*s,y+.3*s],TEAL,2)
    for xx in [-.28,0,.28]: line(ax,[x+xx*s,x+xx*s],[y-.15*s,y+.23*s],WHITE,2)
    line(ax,[x-.4*s,x+.4*s],[y-.24*s,y-.24*s],WHITE,2)
    for xx in [-.27,-.09,.09,.27]:
        ax.add_patch(Circle((x+xx*s,y-.43*s),.035*s,color=PINK))
def graph(ax,x,y,w,h,color=PINK,kind='curve'):
    line(ax,[x,x,x+w],[y+h,y,y],MUTED,1.3)
    t=np.linspace(0,1,100)
    z=.14+.72*(t-.25)**2 if kind=='curve' else .14+.65*t
    line(ax,x+.1*w+.8*w*t,y+h*z,color,4)

# 1: aim.
f,ax=base(1,'One causal curve. Across sites.')
center=(5.05,4.48)
for i in range(10):
    angle=2*math.pi*i/10 + math.pi/10
    x=center[0]+3.55*math.cos(angle); y=center[1]+2.45*math.sin(angle)
    arrow(ax,(x+(center[0]-x)*.18,y+(center[1]-y)*.18),(center[0]+(x-center[0])*.27,center[1]+(y-center[1])*.27),EDGE,1.8)
    bank(ax,x,y,.82)
ax.add_patch(Circle(center,1.03,facecolor=PANEL,edgecolor=PINK,lw=2))
logo=plt.imread(ROOT/'images/flamingo-logo-dark-rounded.png')
ax.imshow(logo,extent=[4.35,5.75,3.78,5.18],zorder=5)
arrow(ax,(6.15,4.48),(10.15,4.48),PINK,3)
card(ax,10.5,2.35,4.65,4.35)
graph(ax,11.1,3.05,3.4,2.8)
text(ax,12.75,6.15,'X → Y',26,ha='center')
text(ax,5.05,1.15,'Data stay local',22,TEAL,ha='center')
text(ax,12.82,1.65,'Non-linear causality',20,ha='center')
notes.append('Aim: estimate a shared non-linear exposure–outcome causal relationship across biobank sites using genetic instruments, without transferring individual records. The diagram is conceptual; ten sites reflect the committed synthetic federation. No formal privacy guarantee is implied by keeping records local. Sources: README.md; writing/methods.md; data/docs/mr-simulation-model.md.')

# 2: methodological map.
f,ax=base(2,'Local computation → shared inference')
card(ax,.7,2.1,4.15,5.15)
text(ax,2.78,6.7,'Each site',21,ha='center')
for x,label in [(1.38,'G'),(2.78,'X'),(4.12,'Y')]:
    ax.add_patch(Circle((x,5.35),.38,facecolor=BG,edgecolor=TEAL if label=='G' else PINK,lw=2))
    text(ax,x,5.35,label,24,ha='center')
arrow(ax,(1.8,5.35),(2.35,5.35)); arrow(ax,(3.2,5.35),(3.7,5.35))
text(ax,2.78,4.15,'U',22,MUTED,ha='center')
arrow(ax,(2.78,4.5),(2.78,4.9),MUTED,1.4); arrow(ax,(3.03,4.3),(3.88,4.95),MUTED,1.4)
bank(ax,2.78,2.95,.7)
for yy,label in [(6.05,'IVW'),(4.5,'FedAvg'),(2.95,'FedMR')]:
    arrow(ax,(4.98,yy),(6.05,yy),PINK)
    card(ax,6.2,yy-.55,3.4,1.1)
    text(ax,6.55,yy,label,20,weight='bold')
    if label=='IVW':
        for j in range(3):
            line(ax,[8.25+j*.12,8.9+j*.12],[yy-.22+j*.22]*2,TEAL,2)
    elif label=='FedAvg':
        for xx,n in [(8.25,2),(8.65,3),(9.05,2)]:
            for k in range(n): ax.add_patch(Circle((xx,yy+(k-(n-1)/2)*.25),.055,color=TEAL))
        for ya in [-.125,.125]:
            for yb in [-.25,0,.25]: line(ax,[8.25,8.65],[yy+ya,yy+yb],TEAL,.7)
        for ya in [-.25,0,.25]:
            for yb in [-.125,.125]: line(ax,[8.65,9.05],[yy+ya,yy+yb],TEAL,.7)
    else:
        for i in range(3):
            for j in range(3): ax.add_patch(Rectangle((8.25+j*.22,yy-.29+i*.22),.14,.14,facecolor=TEAL,alpha=.45+.2*(i==j)))
    arrow(ax,(9.78,yy),(10.8,yy),PINK)
    graph(ax,11.2,yy-.45,2.65,.95,TEAL if label=='FedMR' else PINK,'line' if label=='IVW' else 'curve')
text(ax,7.9,1.5,'Summaries / updates',18,MUTED,ha='center')
text(ax,12.55,1.5,'Estimate / curve',18,MUTED,ha='center')
notes.append('G: genetic instruments; X: exposure; Y: outcome; U: unmeasured confounding. IVW combines per-SNP summary associations into an average slope. FedAvg trains federated neural estimators (including two-stage residual inclusion, 2SRI) via model updates. FedMR aggregates site-centred sufficient statistics and solves two-stage least squares, reproducing the corresponding pooled estimator for its specified basis. FedMR may require one or two rounds depending on basis, first-stage protocol and covariance choice. The route diagrams are conceptual, not fitted results. Source: writing/methods.md; federated_learning/README.md; dashboard/README.md.')

# 3: saved numerical results, no synthetic performance claims.
f,ax=base(3,'Federated ≈ pooled')
rows=list(csv.DictReader((ROOT/'data/results/fedmr.quadratic.csv').open()))
get=lambda name: next(r for r in rows if r['estimator']==name)
fed=get('FedMR quadratic'); pool=get('concatenated quadratic 2SLS'); ivw=get('per-SNP sumstats IVW')
truth=json.loads((ROOT/'data/results/fedmr.quadratic.truth.json').read_text())
p=f.add_axes([.09,.22,.57,.56],facecolor=BG)
x=np.linspace(-3,3,300)
p.plot(x,float(ivw['theta1'])*x,color=MUTED,lw=2.5,ls=':',label='IVW')
p.plot(x,truth['theta1']*x+truth['theta2']*x*x,color=WHITE,lw=3,ls='--',label='Truth')
p.plot(x,float(pool['theta1'])*x+float(pool['theta2'])*x*x,color=TEAL,lw=7,alpha=.65,label='Pooled')
p.plot(x,float(fed['theta1'])*x+float(fed['theta2'])*x*x,color=PINK,lw=2.8,label='FedMR')
p.spines[['top','right']].set_visible(False)
for s in ['left','bottom']: p.spines[s].set_color(EDGE)
p.tick_params(colors=MUTED,labelsize=13); p.set_xlabel('Exposure',color=MUTED,labelpad=10); p.set_ylabel('Causal effect',color=MUTED,labelpad=10)
p.grid(color=EDGE,alpha=.4,lw=.7); p.legend(loc='upper left',frameon=False,fontsize=13,labelcolor=WHITE)
card(ax,11.1,4.48,4.05,2.2)
text(ax,13.12,5.85,'< 10⁻¹⁵',35,TEAL,ha='center',weight='bold')
text(ax,13.12,5.03,'Max |Δθ|',17,MUTED,ha='center')
text(ax,13.12,3.5,'10 sites',24,ha='center')
text(ax,13.12,2.9,'55,182 people',21,ha='center')
text(ax,8,1.03,'Quadratic simulation',16,MUTED,ha='center')
notes.append(f'Real saved quadratic simulation results, not clinical data. Curves show the zero-intercept structural component θ1*x + θ2*x², with truth θ1=0.30, θ2=0.15. FedMR quadratic: θ1={fed["theta1"]}, θ2={fed["theta2"]}. Pooled quadratic: θ1={pool["theta1"]}, θ2={pool["theta2"]}. The maximum coefficient difference reported in the CSV is {fed["abs_diff_from_pooled"]}. This is numerical identity to the corresponding pooled estimator, not error versus truth. IVW represents a single slope and does not recover curvature. N=55,182 across ten sites per writing/results.md and the committed quadratic manifest. Curves do not display uncertainty. Sources: data/results/fedmr.quadratic.csv; data/results/fedmr.quadratic.truth.json; data/simulated_data/federated/quadratic/manifest.csv; writing/results.md.')

# 4: actual application.
f,ax=base(4,'Explore the federation')
screen=OUT/'assets/dashboard.png'
if not screen.exists(): raise FileNotFoundError('Capture the dashboard into presentation/assets/dashboard.png first.')
shot=Image.open(screen)
# Keep the full screenshot, with no invented UI or results.
w=10.4; h=w*shot.height/shot.width
ax.imshow(shot,extent=[4.9,4.9+w,1.05,1.05+h],zorder=3)
for i,(label,y) in enumerate([('Simulate',6.4),('Compare',4.65),('Stress-test',2.9)]):
    ax.add_patch(Circle((1.15,y),.24,facecolor=PINK,edgecolor='none'))
    text(ax,1.15,y,str(i+1),15,BG,ha='center',weight='bold')
    text(ax,1.72,y,label,23,weight='bold')
    if i<2: arrow(ax,(1.15,y-.43),(1.15,y-1.28),EDGE,2)
notes.append('Actual screenshot of the local FLAMINGO Streamlit dashboard from this branch, showing the Sensitivity & invariance area using committed simulation data. Three workflows: simulate heterogeneous biobanks; compare conventional MR, federated learning and non-linear estimators; stress-test assumptions using site exclusion, regularisation, robustness, invariance and what-if perturbations. No new experiment was run for this screenshot. Source: dashboard/app.py; dashboard/sensitivity_tab.py; dashboard/README.md. Launch: cd dashboard && uv run streamlit run app.py.')

with PdfPages(OUT/'FLAMINGO.pdf') as pdf:
    for i,f in enumerate(figures,1):
        f.savefig(OUT/f'assets/slide-{i}.png',dpi=160,facecolor=BG)
        f.savefig(OUT/f'assets/slide-{i}.svg',facecolor=BG)
        pdf.savefig(f,facecolor=BG)
prs=Presentation(); prs.slide_width=Inches(16); prs.slide_height=Inches(9)
for i,note in enumerate(notes,1):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(OUT/f'assets/slide-{i}.png'),0,0,width=prs.slide_width,height=prs.slide_height)
    slide.notes_slide.notes_text_frame.text=note
prs.save(OUT/'FLAMINGO.pptx')
(OUT/'speaker-notes.md').write_text('\n\n'.join(f'## {i}\n\n{n}' for i,n in enumerate(notes,1))+'\n')
# Browser-based deck, all assets local; supports keyboard and fullscreen.
(OUT/'index.html').write_text('''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>FLAMINGO</title><style>html,body{margin:0;width:100%;height:100%;background:#19161d;overflow:hidden}img{width:100vw;height:100vh;object-fit:contain}button{position:fixed;bottom:12px;background:#26212c;color:#fff5f8;border:1px solid #49384d;border-radius:8px;padding:7px 12px;cursor:pointer;opacity:.3}button:hover,button:focus{opacity:1}#prev{left:12px}#next{right:12px}</style><img id="slide" src="assets/slide-1.svg" alt="Slide 1: project aim"><button id="prev" aria-label="Previous slide">←</button><button id="next" aria-label="Next slide">→</button><script>let n=1;const names=['project aim','methods','results','dashboard'];function go(d){n=Math.max(1,Math.min(4,n+d));slide.src='assets/slide-'+n+'.svg';slide.alt='Slide '+n+': '+names[n-1];location.hash=n;}prev.onclick=()=>go(-1);next.onclick=()=>go(1);onkeydown=e=>{if(['ArrowRight',' ','PageDown'].includes(e.key)){e.preventDefault();go(1)}if(['ArrowLeft','PageUp'].includes(e.key)){e.preventDefault();go(-1)}if(e.key==='f')document.documentElement.requestFullscreen();};if(+location.hash.slice(1)>=1&&+location.hash.slice(1)<=4){n=+location.hash.slice(1);go(0)}</script></html>''')
# Contact sheet for inspection.
thumbs=[]
for i in range(1,5):
    im=Image.open(OUT/f'assets/slide-{i}.png').convert('RGB'); im.thumbnail((960,540)); thumbs.append(im)
contact=Image.new('RGB',(1920,1080),BG)
for i,im in enumerate(thumbs): contact.paste(im,((i%2)*960,(i//2)*540))
contact.save(OUT/'preview.png')
print('Created four slides: PPTX, PDF, HTML, SVGs, speaker notes and preview.')
